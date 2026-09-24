#!/usr/bin/env python3
"""Build sweden_daily.csv, sweden_weekly.csv, sweden_regions.csv (population part) and checks.txt
from files in ../raw (run fetch_counts.py first). No numbers are typed in by hand.
"""
import datetime as dt, glob, json, os, re
import openpyxl, pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..')
RAW = os.path.join(ROOT, 'raw')

REGIONS = [  # SCB län code, canonical name
    ('01', 'Stockholm'), ('03', 'Uppsala'), ('04', 'Södermanland'), ('05', 'Östergötland'),
    ('06', 'Jönköping'), ('07', 'Kronoberg'), ('08', 'Kalmar'), ('09', 'Gotland'), ('10', 'Blekinge'),
    ('12', 'Skåne'), ('13', 'Halland'), ('14', 'Västra Götaland'), ('17', 'Värmland'), ('18', 'Örebro'),
    ('19', 'Västmanland'), ('20', 'Dalarna'), ('21', 'Gävleborg'), ('22', 'Västernorrland'),
    ('23', 'Jämtland'), ('24', 'Västerbotten'), ('25', 'Norrbotten')]
NAME = dict(REGIONS)
ALIAS = {'sörmland': '04', 'södermanland': '04', 'jämtland härjedalen': '23', 'jämtland_härjedalen': '23',
         'jämtland': '23', 'västra götaland': '14', 'västra_götaland': '14', 'v. götaland': '14',
         'vgötaland': '14', 'gotlands': '09'}
for c, n in REGIONS:
    ALIAS.setdefault(n.lower(), c)


def code(name):
    k = str(name).strip().lower()
    return ALIAS.get(k) or ALIAS.get(k.replace('_', ' '))


checks = []


def note(s):
    print(s)
    checks.append(s)


def px(t):
    d = json.load(open(os.path.join(RAW, 'fhm_px', f'{t}.px.data.json')))
    cols = [c['code'] for c in d['columns']]
    rows = []
    for r in d['data']:
        v = r['values'][0]
        rows.append(r['key'] + [0 if v == '-' else (None if v in ('..', '.', '') else float(v))])
    return pd.DataFrame(rows, columns=cols[:-1] + ['value'])


# ---------------- population (SCB) ----------------
scb = json.load(open(os.path.join(RAW, 'scb_population_lan_2019_2020.json')))
pop = {}
for r in scb['data']:
    pop.setdefault(r['key'][0], {})[r['key'][1]] = int(r['values'][0])
note(f"SCB population sum 2019-12-31: {sum(p['2019'] for p in pop.values())}, 2020-12-31: {sum(p['2020'] for p in pop.values())}")

# ---------------- FHM daily cases by region (Folkhälsodata acov19DAG) ----------------
c = px('acov19DAG')
c = c[c['Mått'] == '1']
c = c.rename(columns={'Region': 'region_code', 'Dag': 'date', 'value': 'cases_fhm'})[['date', 'region_code', 'cases_fhm']]
c['region_code'] = c['region_code'].replace({'00': 'SE'})

# ---------------- FHM national daily ICU / deaths (xcov19ivavDAG) ----------------
x = px('xcov19ivavDAG')
meta = json.load(open(os.path.join(RAW, 'fhm_px', 'xcov19ivavDAG.px.meta.json')))
mt = dict(zip(meta['variables'][0]['values'], meta['variables'][0]['valueTexts']))
note(f'xcov19ivavDAG measures: {mt}')
x['m'] = x['Mått'].map(lambda k: 'icu_fhm_by_start_date' if 'intensiv' in mt[k] else 'deaths_fhm_by_death_date')
xn = x.pivot_table(index='Dag', columns='m', values='value', aggfunc='sum').reset_index().rename(columns={'Dag': 'date'})
xn['region_code'] = 'SE'

# ---------------- SIR ICU admissions per day, region of ICU unit ----------------
sir = []
for f in glob.glob(os.path.join(RAW, 'sir', 'vtfstart_*.json')):
    reg = os.path.basename(f)[9:-5]
    rc = 'SE' if reg == 'riket' else code(reg)
    assert rc, reg
    d = json.load(open(f))
    for date, n_ep, n_pers in d['DetailedTable']['Rows']:
        sir.append((date, rc, int(n_ep), int(n_pers)))
sir = pd.DataFrame(sir, columns=['date', 'region_code', 'icu_admissions_sir_episodes', 'icu_admissions_sir_persons'])
sir = sir[sir.date.str.match(r'^\d{4}-\d\d-\d\d$')]
tot_reg = sir[sir.region_code != 'SE'].icu_admissions_sir_episodes.sum()
tot_nat = sir[sir.region_code == 'SE'].icu_admissions_sir_episodes.sum()
note(f'SIR 2020 ICU episodes: sum of 21 regions={tot_reg}, national query={tot_nat}')

# ---------------- FHM daily archive -> report-date increments by region ----------------
arch = []
vint_check = []
for f in sorted(glob.glob(os.path.join(RAW, 'fhm_archive', 'Folkhalsomyndigheten_Covid19_*.xlsx'))):
    pub = re.search(r'(\d{4}-\d\d-\d\d)', f).group(1)
    wb = openpyxl.load_workbook(f, read_only=True)
    fo = [s for s in wb.sheetnames if s.startswith('FOHM')]
    if fo:  # sanity: sheet 'FOHM 15 Apr 2020' should match the file name
        sd = dt.datetime.strptime(' '.join(fo[0].split()[1:]), '%d %b %Y').date().isoformat()
        if sd != pub:
            note(f'archive file {pub}: FOHM sheet says {sd}')
    for r in wb['Totalt antal per region'].iter_rows(min_row=2, values_only=True):
        if r[0] is None or str(r[0]).strip() == '':
            continue
        rc = code(r[0])
        assert rc, r[0]
        arch.append((pub, rc, r[1], r[3], r[4]))
    # internal agreement: sum of regional cumulative deaths vs the vintage's own national daily-deaths sheet
    if 'Antal avlidna per dag' in wb.sheetnames:
        nd = sum((r[1] or 0) for r in wb['Antal avlidna per dag'].iter_rows(min_row=2, values_only=True)
                 if isinstance(r[1], (int, float)))
        rd = sum((r[4] or 0) for r in wb['Totalt antal per region'].iter_rows(min_row=2, values_only=True)
                 if r[0] and isinstance(r[4], (int, float)))
        vint_check.append((pub, rd, nd))
vc = pd.DataFrame(vint_check, columns=['pub', 'sum_regions', 'sum_national_daily'])
vc['diff'] = vc.sum_regions - vc.sum_national_daily
note(f'archive internal check (regional cum deaths vs national daily-deaths sheet, same file): '
     f'{(vc["diff"] == 0).sum()} of {len(vc)} vintages agree exactly; max |diff|={vc["diff"].abs().max()}; '
     f'mismatches: {vc[vc["diff"] != 0].to_dict("records")[:10]}')
arch = pd.DataFrame(arch, columns=['pub', 'region_code', 'cum_cases', 'cum_icu', 'cum_deaths'])
arch = arch.sort_values(['region_code', 'pub'])
pubs = sorted(arch.pub.unique())
prev = {p: (pubs[i - 1] if i else None) for i, p in enumerate(pubs)}
arch['prev_pub'] = arch.pub.map(prev)
# defective vintages: national cumulative falls by >20% vs previous file (e.g. 2020-03-27, ICU total = 2)
for k in ['icu', 'deaths']:
    natk = arch.groupby('pub')[f'cum_{k}'].sum()
    bad = [p for i, p in enumerate(pubs) if i and natk[p] < 0.8 * natk[pubs[i - 1]]]
    if bad:
        note(f'archive: cum_{k} treated as missing in defective vintage(s) {bad} '
             f'(national {[int(natk[b]) for b in bad]} vs previous {[int(natk[pubs[pubs.index(b) - 1]]) for b in bad]})')
        arch.loc[arch.pub.isin(bad), f'cum_{k}'] = None
for k in ['cases', 'icu', 'deaths']:
    arch[f'{k}_fhm_reported'] = arch.dropna(subset=[f'cum_{k}']).groupby('region_code')[f'cum_{k}'].diff()
arch['days_since_prev_report'] = [(dt.date.fromisoformat(a) - dt.date.fromisoformat(b)).days if b else None
                                  for a, b in zip(arch.pub, arch.prev_pub)]
neg = arch[(arch.deaths_fhm_reported < 0) | (arch.icu_fhm_reported < 0)]
note(f'archive: {len(pubs)} publication files {pubs[0]}..{pubs[-1]}; negative daily increments (revisions): {len(neg)} region-days')
nat_arch = arch.groupby('pub')[['cum_deaths', 'cum_icu', 'cum_cases']].sum()
ar = arch.rename(columns={'pub': 'date'})[['date', 'region_code', 'cum_deaths', 'deaths_fhm_reported', 'cum_icu',
                                            'icu_fhm_reported', 'days_since_prev_report']]
ar = ar.rename(columns={'cum_deaths': 'cum_deaths_fhm_as_published', 'cum_icu': 'cum_icu_fhm_as_published'})
arn = ar.groupby('date')[['cum_deaths_fhm_as_published', 'deaths_fhm_reported', 'cum_icu_fhm_as_published',
                           'icu_fhm_reported']].sum(min_count=1).reset_index()
arn['days_since_prev_report'] = ar.groupby('date').days_since_prev_report.first().values
arn['region_code'] = 'SE'
ar = pd.concat([ar, arn])

# ---------------- Socialstyrelsen deaths per day of death (national, 2022 vintage) ----------------
wb = openpyxl.load_workbook(os.path.join(RAW, 'sos_avlidna_wayback20220714.xlsx'), read_only=True, data_only=True)
sos = []
for r in wb['Dödsdag'].iter_rows(min_row=9, values_only=True):
    if r[0] and re.match(r'^\d{4}-\d\d-\d\d', str(r[0])):
        v = r[1]
        sos.append((str(r[0])[:10], None if v in ('X', None) else int(v)))
sos = pd.DataFrame(sos, columns=['date', 'deaths_sos_by_death_date'])
sos['region_code'] = 'SE'
first_sos = [str(r[0]) for r in wb['Dödsdag'].iter_rows(min_row=9, max_row=9, values_only=True)][0]
note(f'SoS Dödsdag first row label: {first_sos!r} (the asterisk row pools all deaths up to that date)')

# ---------------- merge daily ----------------
key = ['date', 'region_code']
daily = c.merge(sir, on=key, how='outer').merge(ar, on=key, how='outer').merge(xn, on=key, how='outer') \
    .merge(sos, on=key, how='outer')
daily = daily[(daily.date >= '2020-01-01') & (daily.date <= '2020-12-31')]
# fill structural zeros for SIR where the region query returned no row for the day (SIR omits zero days)
for col in ['icu_admissions_sir_episodes', 'icu_admissions_sir_persons']:
    daily[col] = daily[col].fillna(0).astype(int)
daily['region_name'] = daily.region_code.map(lambda k: 'Sweden' if k == 'SE' else NAME[k])
cols = ['date', 'region_code', 'region_name', 'cases_fhm', 'icu_admissions_sir_episodes',
        'icu_admissions_sir_persons', 'deaths_fhm_reported', 'icu_fhm_reported', 'days_since_prev_report',
        'cum_deaths_fhm_as_published', 'cum_icu_fhm_as_published', 'deaths_fhm_by_death_date',
        'icu_fhm_by_start_date', 'deaths_sos_by_death_date']
# complete date x region grid
grid = pd.MultiIndex.from_product([pd.date_range('2020-01-31', '2020-12-31').strftime('%Y-%m-%d'),
                                   ['SE'] + [k for k, _ in REGIONS]], names=key).to_frame(index=False)
daily = grid.merge(daily, on=key, how='left')
daily['region_name'] = daily.region_code.map(lambda k: 'Sweden' if k == 'SE' else NAME[k])
for col in ['icu_admissions_sir_episodes', 'icu_admissions_sir_persons']:
    daily[col] = daily[col].fillna(0).astype(int)
daily.loc[daily.region_code != 'SE', ['deaths_fhm_by_death_date', 'icu_fhm_by_start_date', 'deaths_sos_by_death_date']] = None
daily = daily[cols].sort_values(['region_code', 'date'])
for col in cols[3:]:
    if col not in ('icu_admissions_sir_episodes', 'icu_admissions_sir_persons'):
        daily[col] = daily[col].astype('Int64')
daily.to_csv(os.path.join(ROOT, 'sweden_daily.csv'), index=False)
note(f'sweden_daily.csv rows={len(daily)}')

# ---------------- weekly: FHM ecov19ivavtid + Socialstyrelsen inpatient admissions ----------------
w = px('ecov19ivavtid')
meta = json.load(open(os.path.join(RAW, 'fhm_px', 'ecov19ivavtid.px.meta.json')))
mt = dict(zip(meta['variables'][1]['values'], meta['variables'][1]['valueTexts']))
keep = {k: v for k, v in mt.items() if 'per 100' not in v}
note(f'ecov19ivavtid measures kept: {keep}')
lab = {k: ('cases_fhm' if v == 'Antal fall' else 'icu_new_fhm' if 'intensiv' in v else 'deaths_fhm') for k, v in keep.items()}
w = w[w['Mått'].isin(keep)]
w['m'] = w['Mått'].map(lab)
w = w.pivot_table(index=['Region', 'År och vecka'], columns='m', values='value', aggfunc='sum').reset_index()
w = w.rename(columns={'Region': 'region_code', 'År och vecka': 'isoweek'})
w['region_code'] = w.region_code.replace({'00': 'SE'})
w = w[w.isoweek.str.startswith('2020') | w.isoweek.isin(['2021W01', '2021W02'])]

# Socialstyrelsen inpatient: 'Inskrivna i slutenvård' (first admission per patient, week x reporting region)
wb = openpyxl.load_workbook(os.path.join(RAW, 'sos_statistik-covid19-inskrivna.xlsx'), read_only=True, data_only=True)
rows = list(wb['Inskrivna i slutenvård'].iter_rows(values_only=True))
hdr_i = [i for i, r in enumerate(rows) if r and any(isinstance(x, str) and x.startswith('vecka ') for x in r)][0]
hdr = rows[hdr_i]
hosp = []
for r in rows[hdr_i + 2:]:
    if not r or r[0] is None:
        continue
    lbl = str(r[0]).strip()
    rc = 'SE' if lbl.startswith('Totalt') else code(lbl)
    if rc is None:
        continue
    for j, h in enumerate(hdr):
        if isinstance(h, str) and h.startswith('vecka '):
            m = re.match(r'vecka (\d+) (\d{4})', h)
            v = r[j]
            hosp.append((rc, f'{m.group(2)}W{int(m.group(1)):02d}', None if v in (None, 'X', '..', '') else v))
hosp = pd.DataFrame(hosp, columns=['region_code', 'isoweek', 'hosp_admissions_sos'])
note(f"SoS inskrivna regions parsed: {sorted(hosp.region_code.unique())}")
w = w.merge(hosp, on=['region_code', 'isoweek'], how='outer')
w = w[w.isoweek.str.startswith('2020') | w.isoweek.isin(['2021W01', '2021W02'])]
w['region_name'] = w.region_code.map(lambda k: 'Sweden' if k == 'SE' else NAME[k])
w = w[['isoweek', 'region_code', 'region_name', 'cases_fhm', 'icu_new_fhm', 'deaths_fhm', 'hosp_admissions_sos']]
for col in ['cases_fhm', 'icu_new_fhm', 'deaths_fhm', 'hosp_admissions_sos']:
    w[col] = pd.to_numeric(w[col]).astype('Int64')
w.sort_values(['region_code', 'isoweek']).to_csv(os.path.join(ROOT, 'sweden_weekly.csv'), index=False)

# ---------------- cross-checks ----------------
# (a) FHM regional weekly deaths summed vs FHM national daily deaths by date of death, per ISO week
xd = xn.copy()
xd['isoweek'] = pd.to_datetime(xd.date).dt.isocalendar().apply(lambda r: f'{r.year}W{r.week:02d}', axis=1)
wk_nat = xd.groupby('isoweek')[['deaths_fhm_by_death_date', 'icu_fhm_by_start_date']].sum()
wr = w[w.region_code != 'SE'].groupby('isoweek')[['deaths_fhm', 'icu_new_fhm']].sum()
cmp_ = wr.join(wk_nat, how='inner').loc['2020W10':'2020W30']
note('Weekly: sum over regions of FHM weekly deaths/ICU vs FHM national daily-by-date summed to ISO week:\n'
     + cmp_.to_string())
# (b) archive vintages vs final: national cumulative deaths as published on dates vs final death-date series
fin = xn.set_index('date').deaths_fhm_by_death_date.cumsum()
for d in ['2020-04-15', '2020-05-15', '2020-06-15', '2020-09-15', '2020-12-30']:
    if d in nat_arch.index:
        note(f'as published {d}: cum deaths={int(nat_arch.loc[d, "cum_deaths"])}, cum ICU={int(nat_arch.loc[d, "cum_icu"])};'
             f' final FHM deaths with death date <= {d}: {int(fin.loc[:d].iloc[-1])}')
# (c) totals for 2020
nat = daily[daily.region_code == 'SE']
note('2020 totals, national: ' + ', '.join(f'{k}={nat[k].sum()}' for k in
     ['cases_fhm', 'icu_admissions_sir_episodes', 'icu_admissions_sir_persons', 'deaths_fhm_by_death_date',
      'icu_fhm_by_start_date', 'deaths_sos_by_death_date', 'deaths_fhm_reported']))
reg = daily[daily.region_code != 'SE']
note('2020 totals, sum of regions: ' + ', '.join(f'{k}={reg[k].sum()}' for k in
     ['cases_fhm', 'icu_admissions_sir_episodes', 'deaths_fhm_reported', 'icu_fhm_reported']))
# final regional totals (FHM archived last 2020 vintage) vs weekly FHM sums
last = arch[arch.pub == max(p for p in pubs if p <= '2020-12-31')].set_index('region_code')
wk20 = w[(w.region_code != 'SE') & w.isoweek.str.startswith('2020')].groupby('region_code')[['deaths_fhm', 'icu_new_fhm']].sum()
note('Per region: cum deaths as published last 2020 file vs sum of current weekly 2020 deaths:\n'
     + last[['cum_deaths', 'cum_icu']].join(wk20).to_string())

# ---------------- regions (population part; sero merged later by merge_regions.py) ----------------
reg_rows = [{'region_code': k, 'region_name': n, 'population_2019_12_31': pop[k]['2019'],
             'population_2020_12_31': pop[k]['2020']} for k, n in REGIONS]
pd.DataFrame(reg_rows).to_csv(os.path.join(ROOT, 'raw', 'regions_population.csv'), index=False)
open(os.path.join(ROOT, 'checks.txt'), 'w').write('\n\n'.join(checks) + '\n')

# (d) final FHM Excel (Wayback, 23 Mar 2023) vs current Folkhälsodata PxWeb
wb = openpyxl.load_workbook(os.path.join(RAW, 'Folkhalsomyndigheten_Covid19_wayback20230323.xlsx'), read_only=True)
xl = {}
for r in wb['Antal avlidna per dag'].iter_rows(min_row=2, values_only=True):
    if isinstance(r[0], dt.datetime) and r[0].year == 2020:
        xl[r[0].date().isoformat()] = r[1]
px_d = xn.set_index('date').deaths_fhm_by_death_date
dd = [(d, v, px_d.get(d)) for d, v in xl.items() if px_d.get(d) != v]
note(f'FHM Excel 2023 vs PxWeb daily national deaths 2020: Excel sum={sum(xl.values())}, '
     f'PxWeb sum={int(px_d.loc["2020-01-01":"2020-12-31"].sum())}, days differing={len(dd)} {dd[:8]}')
vr = {}
for r in wb['Veckodata Region'].iter_rows(min_row=2, values_only=True):
    if r[0] == 2020:
        vr[code(r[2])] = vr.get(code(r[2]), 0) + (r[7] or 0)
note('FHM Excel 2023 Veckodata Region 2020 deaths by region vs PxWeb weekly: '
     + str({k: (vr[k], int(wk20.loc[k, 'deaths_fhm'])) for k in sorted(vr) if vr[k] != wk20.loc[k, 'deaths_fhm']} or 'all equal'))
open(os.path.join(ROOT, 'checks.txt'), 'w').write('\n\n'.join(checks) + '\n')
